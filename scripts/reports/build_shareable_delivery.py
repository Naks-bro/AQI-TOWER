"""Two audience bundles from existing R03M/IH02 evidence, not a release.
Requires ReportLab, Pillow, openpyxl, pypdf, pypdfium2; use bundled runtime.
"""
import csv
import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape
from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.colors import HexColor, white
from reportlab.graphics.shapes import Drawing, Circle, Rect, Line, String
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter
import pypdfium2 as pdfium

R = Path(__file__).resolve().parents[2]
O = R/'reports/shareable_20261005'
V = O/'visuals'
IH = R/'reports/prototype_d01/internal_handoff_20261005'
M = R/'reports/prototype_d01/mechanical_package'
S = O/'stakeholder'
T = O/'technical'
for d in (S,T,O/'qa'):
    d.mkdir(parents=True,exist_ok=True)
mesh = json.loads((M/'cad_mesh.json').read_text())
pressure = json.loads((IH/'CURRENT_PRESSURE_BUDGET.json').read_text())
central = next(x for x in pressure['rows'] if x['total_m3h']==300 and x['reducer_K']==1)
assert abs(central['filter_allowance_Pa']-14.316071817116585)<1e-8
assert len(mesh)==493
assert hashlib.sha256((M/'D01_R03M_ASSEMBLY.FCStd').read_bytes()).hexdigest()=='fb28a31315968509dcdbbb2caac8c6d69c84ecb94a35ccd75937a5d9ecaa33a6'

# Self-contained offline viewer; actual mesh coordinates rounded only for display.
display_mesh=[{k:p[k] for k in ('name','kind','explode','note','triangles')} |
              {'vertices':[[round(x,3) for x in v] for v in p['vertices']]} for p in mesh]
html=(R/'threejs/presentation.template.html').read_text(encoding='utf-8')
html=html.replace('__CAD_DATA__',json.dumps(display_mesh,separators=(',',':')).replace('</','<\\/'))
html=html.replace('__VIEWER_CODE__',(V/'viewer.js').read_text(encoding='utf-8'))
(O/'AQI_TOWER_3D_REVIEW.html').write_text(html,encoding='utf-8')
shutil.copy2(R/'threejs/node_modules/three/LICENSE',O/'THREE_JS_LICENSE.txt')

fontpath=Path('C:/Windows/Fonts/arial.ttf')
font=ImageFont.truetype(str(fontpath),14)
for directory,output,label in [('frames','assembly_exploded.gif','CAD assembly explanation | R03M | NOT build-approved'),
                              ('air_frames','illustrative_airflow.gif','ILLUSTRATIVE AIR PATHS | NOT CFD or measured performance')]:
    frames=[]
    for path in sorted((V/directory).glob('*.png')):
        im=PILImage.open(path).convert('RGB')
        draw=ImageDraw.Draw(im)
        draw.rectangle((0,0,im.width,40),fill='#173d49')
        draw.text((12,11),label,fill='white',font=font)
        frames.append(im)
    assert len(frames)==24, f'Incomplete Blender render: {directory}'
    frames[0].save(V/output,save_all=True,append_images=frames[1:],duration=167,loop=0)

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyAQI',fontName='Helvetica',fontSize=11,leading=16,spaceAfter=10,textColor=HexColor('#263d48')))
styles.add(ParagraphStyle(name='SmallAQI',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=8,textColor=HexColor('#536b73')))
styles['Title'].fontSize=30;styles['Title'].leading=35;styles['Title'].textColor=HexColor('#174b58')
styles['Heading1'].fontSize=21;styles['Heading1'].leading=26;styles['Heading1'].textColor=HexColor('#174b58')
styles['Heading2'].fontSize=13;styles['Heading2'].leading=17;styles['Heading2'].textColor=HexColor('#174b58')
def p(t,small=False):return Paragraph(t,styles['SmallAQI' if small else 'BodyAQI'])
def h(t):return Paragraph(t,styles['Heading1'])
def sub(t):return Paragraph(t,styles['Heading2'])
def picture(path,width=460):
    im=PILImage.open(path)
    return Image(str(path),width=width,height=width*im.height/im.width)
def table(header,rows,widths):
    tab=Table([[p(escape(str(x)),True) for x in row] for row in [header]+rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
    tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#d8eceb')),('GRID',(0,0),(-1,-1),.4,HexColor('#d1dcdf')),('ROWBACKGROUNDS',(0,1),(-1,-1),[white,HexColor('#f1f6f7')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    return tab
def footer(c,d):
    c.setFillColor(HexColor('#174b58'));c.rect(0,0,595,31,fill=True,stroke=False)
    c.setFillColor(white);c.setFont('Helvetica',8)
    c.drawString(36,12,'AQI Tower | 5 October 2026 | Digital review only | Physical prototypes 0')
    c.drawRightString(557,12,str(d.page))
def pdf(name,story):
    SimpleDocTemplate(str(name),pagesize=(595,842),leftMargin=38,rightMargin=38,topMargin=34,bottomMargin=46).build(story,onFirstPage=footer,onLaterPages=footer)
def bubbles():
    d=Drawing(515,165)
    for i in range(3):
        x=93+i*164
        d.add(Circle(x,90,61,fillColor=HexColor('#e8f5f1'),strokeColor=HexColor('#69aaa4'),strokeDashArray=[4,4]))
        d.add(Rect(x-12,62,24,58,fillColor=HexColor('#246f7a'),strokeColor=None))
        d.add(String(x,43,'Future tower',textAnchor='middle',fontSize=10,fillColor=HexColor('#174b58')))
        if i<2:d.add(Line(x+65,90,x+98,90,strokeColor=HexColor('#69aaa4'),strokeWidth=2))
    d.add(String(257,8,'Concept only. No measured radius, spacing or guaranteed protected zone.',textAnchor='middle',fontSize=9))
    return d

six=[('Vision for 2026','Build and safely test one indoor prototype. Use those results to decide whether a larger outdoor tower is worth developing.'),
     ('Work completed','An integrated CAD model, mechanical drawings, an electrical reference plan, calculations and physical-test instructions are ready for review. No physical unit has been built.'),
     ('Key decisions','Start with particle filtration indoors. Use a custom housing with two parallel filters and four fans. Keep water, gas-filter media and solar outside this first build.'),
     ('Unresolved challenges','Actual filter resistance, sealing, guard strength, cabinet stability and electrical protection still need closure. Outdoor coverage is unknown.'),
     ('Support requirements','Fund exact parts after approval, qualified design review, local fabrication, safe assembly and rented or paid measurement services. We do not assume a college lab.'),
     ('Plan for upcoming month','Assign reviewers, resolve supplier inputs, agree test targets and collect quotes. Release and build only after approval. Choose a final project name. Dates depend on responses and workshop capacity.')]
stake=[Paragraph('AQI Tower',styles['Title']),p('Prototype development and funding review'),
       p('We are preparing a small air-cleaning prototype. We need support to finish the engineering review, build it safely and measure whether it works.'),picture(V/'assembly.png',470),
       p('Actual CAD-derived rendering, not a physical product photo. Purple parts reserve control-box and cable space. Guard openings are not represented accurately.',True),
       p('<b>Current position:</b> digital design exists; construction and performance are not yet approved. AQI Tower is a working name.'),PageBreak(),
       h('Report 1 Executive Summary'),p('A short update for the college, NGO and funding team.')]
for title,body in six:stake += [sub(title),p(body)]
stake += [PageBreak(),h('Report 2 Comprehensive Report'),sub('Vision for 2026'),p(six[0][1]),
          sub('Project idea and the clean air bubble'),p('The long-term idea is a larger outdoor cylindrical tower. Each tower would aim to improve air near people. Several towers might form a chain of local cleaner-air zones.'),bubbles(),
          p('A bubble is a simple way to explain the goal. It is not a sealed dome or a proven clean-air boundary. Outdoor wind mixes air continuously, so we cannot promise a bubble size or cleaning percentage.'),
          p('The current prototype is a small rectangular indoor unit, not the final outdoor tower. Keeping the first build simple helps us inspect, transport, measure and repair it.'),
          sub('Where it may be used later'),p('Potential sites include sheltered waiting areas and other managed public locations. No site has been selected. Fully outdoor use needs separate heat, rain, dust, corrosion, anchoring and vandal-resistance design.'),PageBreak(),
          h('How the prototype works'),picture(V/'cutaway.png',415),
          table(['Part','Simple explanation'],[['Two filters','Air enters through two opposite filter branches. They work side by side, not in sequence.'],['Four fans','Pull air into a shared chamber and discharge it through the top.'],['Cabinet seals and guards','Keep the parts supported, limit leaks and restrict access to moving fans. These details still need approval.']],[135,384]),
          p('Cutaway view removes panels and guards for explanation. Never operate an unguarded unit. Colours identify parts; they do not show measured pollution, flow speed or efficiency.',True),PageBreak(),
          h('Work completed'),picture(V/'exploded.png',390),
          p('The current model coordinates the cabinet, filters, seals, fan plate, guards and proposed hardware. Drawings, calculations, review sheets and test software are available for the internal team.'),
          sub('What the computer work tells us'),p('We can check nominal fit and explore pressure losses. We cannot yet state actual airflow, cleaning efficiency, noise, power or filter lifetime. The fan and filter must work together under load.'),
          sub('Important evidence boundary'),p('Earlier coloured CFD studies belong to a different large-tower geometry. They were partial or nonconverged and do not validate this prototype. The new animation explains assembly and the intended air path; it is not a new CFD result.'),PageBreak(),
          h('Key decisions and unresolved challenges'),sub('Key decisions'),p(six[2][1]),
          table(['Question','Current answer'],[['Has a tower been physically tested?','No. Physical prototypes built: 0.'],['Are parts and wiring released?','No. The integrated design is for review only.'],['Will it remove gases?','Not demonstrated. The first scope is particle filtration.'],['Does it have a verified HEPA rating?','No whole-device HEPA classification.'],['How large is the outdoor bubble?','Unknown. Outdoor benefit and spacing need field measurements.'],['Are costs and dates confirmed?','No. Quotes, supply times and fabrication capacity are still unknown.']],[235,284]),
          p('The main missing work is component evidence, mechanical and electrical approval, then construction and controlled measurements. Another attractive model does not close those items.'),PageBreak(),
          h('Support requirements and upcoming month'),sub('Support requirements'),p(six[4][1]),
          table(['Funding stage','What funding should enable'],[['Review first','Resolve guard, structure, filter seal and electrical protection with named reviewers.'],['Build after release','Purchase approved parts and outsource the required cabinet and guard fabrication.'],['Measure after commissioning','Rent suitable instruments or commission a test service; record airflow, pressure, particles, noise and power.']],[150,369]),
          sub('Plan for upcoming month'),p(six[5][1]),
          sub('The next measurable milestone'),p('A review-approved prototype assembly followed by a safe first run with recorded airflow and pressure. A later controlled room test should compare background particle decay with purifier-on trials.'),
          p('Supplier responses, fabrication, delivery, qualified commissioning and physical testing are separate tasks. No completion date or performance percentage is promised.'),
          p('Use STAKEHOLDER_FUNDING.xlsx for staged funding discussions. Meetings content is omitted as requested; no meeting count has been invented.',True)]
pdf(S/'AQI_TOWER_STAKEHOLDER_REPORT.pdf',stake)

tech=[Paragraph('AQI Tower Engineering Handoff',styles['Title']),p('Current D01 R03M integration with IH02 evidence'),picture(V/'exploded.png',420),
      p('<b>NOT FOR FABRICATION OR ENERGIZATION.</b> Authoritative manufacturing geometry remains the FreeCAD R03M study; Blender and browser meshes are display derivatives.'),
      p('The following review pages lead into the unchanged 41-page IH02 engineering reference. Return named decisions with evidence in TECHNICAL_REVIEW.xlsx. Do not mix legacy large-tower mains/HEPA assumptions with this 12 V indoor demonstrator.'),PageBreak(),
      h('Current pressure and geometry'),
      table(['Quantity','Current screen or evidence'],[['Total flow','300 m3/h assumed, not measured'],['Parallel branches','4 fans at 75 m3/h each; 2 filters at 150 m3/h each'],['OEM fan head','27.63 Pa at stated curve condition'],['Nonfilter loss','13.31 Pa under K1 reducer scenario'],['Remaining filter head','14.32 Pa per parallel filter branch'],['CAD inventory','493 source objects; tessellated display derived from saved R03M'],['Display bounding envelope','595 x 504 x 631 mm including reservations and projecting hardware; not a fabrication specification']],[175,344]),
      p('The two parallel filters do not add their pressure drops in series. Clean/loaded filter curves, installed/static fan curve compatibility, loss coefficients and branch equality remain unresolved. This screen is not an operating-point prediction.'),
      sub('Colour and animation interpretation'),p('Teal identifies cabinet/guards, amber filters, dark blue fan envelopes, purple reserved control space. Motion shows part separation or hand-authored air paths, not velocity, turbulence, deposition or removal. Guard faces remain solid tessellations, not the proposed perforation geometry.'),PageBreak(),
      h('Review responsibilities and construction path'),
      table(['Discipline','Closure required'],[['Mechanical','Guard strength/reach, seams and attachments; cabinet joints and stability; real filter tolerance, gasket compression, safe service access and whole mass/CG.'],['Electrical','Hazard allocation, independent protective functions, actual DC switching/isolation/protection, wire/connector ratings, cable restraints and enclosure. No guessed terminal assignments.'],['OEM-dependent','Exact fan/extension connector and clamping dimensions; delivered filter/frame properties. Sent enquiries are pending.'],['Commissioning and test','Confirm premises and targets; review assembly, isolation/restart and abnormal behaviour before collecting measured performance.']],[115,404]),
      sub('Critical path'),p('Exact component evidence and accepted targets -> mechanical/electrical review -> revision-specific fabrication release -> supplier delivery and fabrication -> assembly and qualified commissioning -> airflow/pressure and controlled particle tests. These dependencies are not completion dates.'),
      sub('Outdoor boundary'),p('Current fan temperature limits, enclosure and materials do not establish suitability for Indian outdoor heat or public misuse. Weather protection, tamper-resistant service access, anchoring and wind-dependent performance require a separate qualified design and approved site.'),
      sub('Files to use'),p('Native CAD, STEP, 15 DXFs, current tools, measurement templates and source registers are in engineering/. TECHNICAL_REVIEW.xlsx combines parts, release, pressure and hardware sheets. The Blender file is for presentation only.'),PageBreak(),
      h('Historical CFD evidence'),picture(R/'results/FILTER_H13/central_velocity_pressure_CASE_M.png',505),
      p('Existing Phase 9 CASE_M plot: KVO 250 + GEO_C + H13 approximation from a single manufacturer data point. Colours show calculated speed and gauge pressure for that historical case, not R03M performance.'),
      p('GEO_C is an earlier air-path geometry trial. It is not the current manufacturing assembly. Historical CFD is partial/nonconverged; the later numerical bypass seal is not a proven physical seal. No new R03M CFD has been run for this delivery.'),
      p('For the current prototype, use the pressure screen and obtain measured filter/installed-flow evidence. The following IH02 appendices are preserved unchanged.')]
pdf(T/'REVIEW_INTRO.pdf',tech)
writer=PdfWriter()
for path in (T/'REVIEW_INTRO.pdf',IH/'03_ENGINEERING_DRAWINGS_AND_REVIEW.pdf'):
    writer.append(str(path))
with (T/'AQI_TOWER_TECHNICAL_HANDOFF.pdf').open('wb') as f:writer.write(f)

def readcsv(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.reader(f))
def sheet(wb,name,rows):
    ws=wb.create_sheet(name)
    for row in rows:ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='174B58');cell.font=Font(color='FFFFFF',bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EEF5F5')
        ws.row_dimensions[row[0].row].height=48
    for i,col in enumerate(ws.columns,1):
        ws.column_dimensions[get_column_letter(i)].width=min(55,max(14,max(len(str(c.value or '')) for c in col[:20])*.65))
    # Wrapped descriptions must remain readable, including long review holds.
    import math
    for row in ws.iter_rows(min_row=2):
        lines=max(sum(max(1,math.ceil(len(line)/(ws.column_dimensions[cell.column_letter].width*.9))) for line in str(cell.value or '').split('\n')) for cell in row)
        ws.row_dimensions[row[0].row].height=max(36,lines*15+8)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    return ws
milestones=[['Milestone','Current state','Evidence needed','Timing'],['Engineering review','Open','Named mechanical/electrical decisions and linked evidence','Dependent on reviewers and OEM inputs'],['Construction release','Not released','Frozen components, drawings and electrical schedule','After review'],['First safe run','Not done','Signed inspection/commissioning and measured airflow/pressure','After delivery, fabrication and assembly'],['Indoor performance','Not done','Controlled repeated trials and raw logs','After safe run'],['Outdoor pilot','Separate future scope','Site permission, weather/public-safety design and wind-dependent tests','Not scheduled']]
claims=[['Claim','Evidence status','Permitted wording'],['Physical prototype','0 built','We have prepared the digital design for review.'],['Flow and cleaning','Unmeasured','The calculations are conditional; tests are planned.'],['Outdoor bubble','Unproven','The chain of zones is a research concept.'],['Costs and stock','Unquoted','Funding will be based on verified quotes.'],['Final name','Not selected','AQI Tower is the working name.']]
for technical,path in [(False,S/'STAKEHOLDER_FUNDING.xlsx'),(True,T/'TECHNICAL_REVIEW.xlsx')]:
    wb=Workbook();wb.remove(wb.active)
    sheet(wb,'Read me',[['Topic','Meaning'],['Revision','D01 R03M / IH02 / 5 October 2026'],['Status','Review only. Not order, fabrication or energization authorization.'],['Prices','Blank means UNKNOWN; no prices have been invented.'],['Evidence','CAD/digital checks are not physical results.'],['Team','College/NGO funding; no college lab assumed; qualified physical services still needed.']])
    sheet(wb,'Milestones',milestones)
    sheet(wb,'Claims',claims)
    rows=readcsv(IH/'QUOTE_AND_FUNDING_REGISTER.csv')
    rows[0]+=['Quoted_line_total_INR']
    for n,row in enumerate(rows[1:],2):
        row += [f'=IF(AND(ISNUMBER(C{n}),ISNUMBER(F{n}),ISNUMBER(G{n}),ISNUMBER(H{n})),C{n}*F{n}+SUM(G{n}:H{n}),"UNQUOTED")']
        try:row[2]=float(row[2])
        except ValueError:pass
    sheet(wb,'Funding quotes',rows)
    if technical:
        for name,file in [('Release decisions','REVIEW_AND_RELEASE_REGISTER.csv'),('Parts register','COMBINED_PARTS_REGISTER.csv')]:sheet(wb,name,readcsv(IH/file))
        keys=list(pressure['rows'][0]);sheet(wb,'Pressure screening',[keys]+[[row[k] for k in keys] for row in pressure['rows']])
        sheet(wb,'Pressure assumptions',[['Field','Value'],['Status',pressure['status']],['Assumptions',pressure['assumptions']],['Filter data','No measured clean or loaded curve; allowance is not actual loss']])
        sheet(wb,'Fasteners',readcsv(M/'FASTENER_SCHEDULE.csv'))
        sheet(wb,'Display dimensions',[['CAD object','Kind','X mm','Y mm','Z mm','Limit']]+[[p['name'],p['kind']]+[round(max(v[i] for v in p['vertices'])-min(v[i] for v in p['vertices']),3) for i in range(3)]+['Mesh bounding box; no tolerance or manufacturing release'] for p in mesh])
    wb.save(path)
    check=load_workbook(path);assert check['Read me']['B2'].value.startswith('D01')

for dest in (S,T):
    (dest/'visuals').mkdir(exist_ok=True)
    for name in ('assembly.png','exploded.png','cutaway.png','assembly_exploded.gif','illustrative_airflow.gif'):
        shutil.copy2(V/name,dest/'visuals'/name)
    shutil.copy2(O/'AQI_TOWER_3D_REVIEW.html',dest/'AQI_TOWER_3D_REVIEW.html')
    shutil.copy2(O/'THREE_JS_LICENSE.txt',dest/'THREE_JS_LICENSE.txt')
shutil.copytree(IH/'engineering',T/'engineering',dirs_exist_ok=True,ignore=shutil.ignore_patterns('__pycache__','*.pyc'))
shutil.copy2(V/'AQI_R03M_REVIEW.blend',T/'AQI_R03M_REVIEW.blend')
(T/'BLENDER_VIEW_GUIDE.md').write_text('''# Blender presentation guide

Open AQI_R03M_REVIEW.blend. Use the active/latest scene named AQI R03M Review Presentation (currently .001); other pre-existing scenes were preserved, not overwritten.

Frame 1: assembled. Frames 30-36: exploded. Frame 48: assembled again. Camera zoom is animated to keep exploded parts inside the view. Space/play controls animate the assembly; render the camera for final lighting.

493 source parts are actual CAD tessellations, in metres for display. Five studio objects and 32 illustrative air markers bring the latest scene to 530 objects. The source CAD remains authoritative; Blender mesh dimensions do not establish fabrication tolerances. Guards remain solid envelopes.

Air markers are hidden from render by default. The supplied air GIF was rendered from a temporary cutaway with CAD animations cleared, cabinet/guard panels hidden and air markers enabled. These paths are hand-authored, NOT CFD. The source file remains the assembly animation. Use the provided GIF or offline viewer for the ready-to-share airflow explanation.

Sources: scripts/visualization/build_blender_review.py and render_blender_review.py. The latter renders a temporary background copy and does not save over the source scene. No paid assets or copied product model are used.
''',encoding='utf-8')
(T/'historical_cfd').mkdir(exist_ok=True)
shutil.copy2(R/'results/FILTER_H13/central_velocity_pressure_CASE_M.png',T/'historical_cfd/PHASE9_CASE_M_NOT_R03M.png')
for name in ('CURRENT_PRESSURE_BUDGET.json','TOOL_USAGE.md','FILTER_CURVE_BLANK.csv'):
    shutil.copy2(IH/name,T/name)
(S/'START_HERE.md').write_text('''# AQI Tower stakeholder bundle

Read AQI_TOWER_STAKEHOLDER_REPORT.pdf first. It separates Report 1 Executive Summary from Report 2 Comprehensive Report. Meetings content is omitted as requested; final project name is still needed.

Open AQI_TOWER_3D_REVIEW.html directly in a WebGL-capable browser. No server or internet is required. Simple view explains the device; Technical view shows the conditional pressure and release holds. Use GIFs in visuals/ for messaging/slides.

STAKEHOLDER_FUNDING.xlsx is the quote and milestone workbook. Prices, stock and completion dates are unknown. Line totals require numeric quantity, unit price, total-line tax and total-line delivery; explicitly enter zero only if the quote confirms zero. Do not treat a quoted line or a rendering as construction approval.

Physical prototypes built: 0. Air animation is illustrative, NOT CFD. Outdoor bubbles are unproven. R03M guards remain solid CAD envelopes; purple objects reserve space. No purchase or contact is performed by this package.
''',encoding='utf-8')
(T/'START_HERE.md').write_text('''# AQI Tower technical bundle

Read AQI_TOWER_TECHNICAL_HANDOFF.pdf: current four-page introduction followed by the preserved 41-page IH02 engineering reference. Record actual decisions in TECHNICAL_REVIEW.xlsx. Not for fabrication, ordering or energization.

Authoritative assembly: engineering/reports/prototype_d01/mechanical_package/D01_R03M_ASSEMBLY.FCStd. STEP, exploded CAD, harness reservation derivative and 15 review DXFs accompany it. Blender and the offline viewer are tessellated presentation derivatives, not manufacturing models or CFD validation.

Run tools from engineering/ using TOOL_USAGE.md. The current pressure budget and blank filter curve are supplied. Supply actual measured data and calibration/source evidence; do not substitute animation dots for results.

Outstanding: exact filter resistance/tolerances, guard/structure/seal release, cable dimensions, hazard/protection/rated electrical enclosure, premises/targets, quotes, fabrication, commissioning and physical tests. OEM technical replies remain pending in reviewed records.

No lab assumed. No physical prototype built. Outdoor heat, rain, vandal resistance, site anchoring and measurable benefit remain separate future engineering.
''',encoding='utf-8')

# Evidence-based GitHub coverage inventory: omit regenerable bulk and credentials.
ignored=subprocess.run(['git','ls-files','--others','--ignored','--exclude-standard'],cwd=R,capture_output=True,text=True,check=True).stdout.splitlines()
groups={}
for name in ignored:
    path=R/name
    inaccessible=False
    try:
        if not path.is_file():continue
        size=path.stat().st_size
    except OSError:
        size=0
        inaccessible=True
    if name.startswith('reports/shareable_20261005/') or name.endswith('.blend1'):category='Reproducible presentation QA frames and Blender save backup'
    elif name.startswith('threejs/node_modules/') or '__pycache__' in name:category='Dependencies and caches'
    elif name.startswith('cfd/'):category='Generated raw solver fields meshes logs'
    elif name.startswith('GIF/'):category='Historical viewer recordings retained locally'
    elif name.startswith('tmp/') or name.endswith('.FCBak'):category='Scratch and recoverable CAD backups'
    elif name.endswith('.zip'):category='Redundant historical ZIP exports; source folders tracked'
    else:category='Other ignored local output; see existing ignore policy'
    counts=groups.setdefault(category,[0,0,0]);counts[0]+=1;counts[1]+=size;counts[2]+=int(inaccessible)
with (O/'GITHUB_COVERAGE.csv').open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(['Retained_local_category','File_count','Readable_bytes_subtotal','Inaccessible_size_count','Reason']);w.writerows([[k,*v,'Not silently deleted or claimed uploaded; inaccessible WSL links excluded from byte subtotal'] for k,v in sorted(groups.items())])
(O/'README.md').write_text('''# Shareable AQI Tower delivery

5 October 2026. Two audiences, one unchanged engineering basis: D01 R03M / IH02.

- stakeholder/: plain-language Executive Summary and Comprehensive Report, funding workbook, offline dual-mode viewer and GIFs.
- technical/: review introduction plus preserved IH02 drawings, workbook, authoritative CAD/STEP/DXFs, current tools and editable Blender presentation.
- visuals/: actual CAD-derived renders and annotated GIFs. Colour identifies parts; air motion is illustrative, not CFD.
- GITHUB_COVERAGE.csv: excluded generated/dependency/history categories retained locally. GitHub is not an exact mirror of raw solver time folders or dependencies.

Extract either ZIP before opening its HTML. WebGL is required for 3D; PDFs and GIFs remain usable without it. The viewer is offline and includes the Three.js MIT licence. Source build tools are scripts/visualization/ and scripts/reports/build_shareable_delivery.py.

No new physical results or build approval. Exact component/filter data, mechanical/electrical release, quotations, site, fabrication and safe commissioning remain outstanding. No fabricated price or completion date. Final device name not selected.
''',encoding='utf-8')
for dest,name in [(S,'AQI_TOWER_STAKEHOLDER_BUNDLE.zip'),(T,'AQI_TOWER_TECHNICAL_BUNDLE.zip')]:
    files=[f for f in dest.rglob('*') if f.is_file() and f.name!='MANIFEST.json']
    manifest={str(f.relative_to(dest)).replace('\\','/'):hashlib.sha256(f.read_bytes()).hexdigest() for f in files}
    (dest/'MANIFEST.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    with zipfile.ZipFile(O/name,'w',zipfile.ZIP_DEFLATED) as z:
        for f in dest.rglob('*'):
            if f.is_file():z.write(f,f.relative_to(dest))
    with zipfile.ZipFile(O/name) as z:
        assert all(hashlib.sha256(z.read(n)).hexdigest()==sha for n,sha in manifest.items())

# Render every new PDF page for inspection; preserved appendices unchanged.
for path in (S/'AQI_TOWER_STAKEHOLDER_REPORT.pdf',T/'REVIEW_INTRO.pdf'):
    doc=pdfium.PdfDocument(str(path))
    for i in range(len(doc)):
        doc[i].render(scale=1.5).to_pil().save(O/'qa'/f'{path.stem}_{i+1}.png')
    print(path.name,'pages',len(doc))
print('Technical combined pages',len(PdfReader(T/'AQI_TOWER_TECHNICAL_HANDOFF.pdf').pages))
print('Shareable bundles built; assertions and ZIP manifests pass.')
