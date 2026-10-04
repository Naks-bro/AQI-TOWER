"""D01 R01 geometry: compare OEM numeric curve with historical brochure reading.
No curve averaging, performance claim, CAD change or procurement release.
"""
import csv
import hashlib
import json
from pathlib import Path
import openpyxl
import d01_pressure as p

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'reports/prototype_d01/component_validation'
SRC = ROOT / 'data/prototype_d01/sources'
URL = 'https://support.arctic.de/products/p14-max/techdocs/P14%20Max%20-%20PQ%20Curve%20for%20CFD.XLSX'

def main():
    OUT.mkdir(exist_ok=True)
    source = SRC / 'ARCTIC_P14_Max_PQ_CFD.xlsx'
    sheet = openpyxl.load_workbook(source, data_only=True)['P14 Max']
    assert sheet['C7'].value == 2800 and 'mmH2O' in sheet['C12'].value
    assert 'ft^3/min' in sheet['B12'].value
    points = [(sheet.cell(r, 2).value, sheet.cell(r, 3).value) for r in range(14, 28)]
    assert all(isinstance(v, (float, int)) for row in points for v in row)
    xs, ys = zip(*points)
    assert all(a < b for a, b in zip(xs, xs[1:]))
    assert all(a > b for a, b in zip(ys, ys[1:]))
    # R01 solid guard rim: 300 mm square active area with assumed 40% porosity.
    p.C['guard_open_fraction'] = .4 * (300 / 320)**2
    def numeric(q):
        return p.interp(q / (4*p.CFM), xs, ys)*p.MMH2O
    def solve(fn, ref):
        a, b = 0., 4 * xs[-1] * p.CFM
        for _ in range(80):
            mid = (a+b)/2
            if fn(mid) > sum(p.budget(mid, ref).values()): a = mid
            else: b = mid
        return (a+b)/2
    curves = {'historical_brochure_digitization': p.fan, 'OEM_numeric_2800rpm': numeric}
    results = []
    for name, fn in curves.items():
        for ref in (5, 10, 20):
            q = solve(fn, ref)
            results.append(dict(curve=name, assumed_filter_Pa_at_total400=ref,
                total_flow_m3h=q, each_filter_flow_m3h=q/2,
                fan_static_Pa=fn(q), loss_breakdown_Pa=p.budget(q, ref),
                closure_error_Pa=fn(q)-sum(p.budget(q, ref).values())))
    headroom = []
    for q in (250, 300, 350, 400):
        other = sum(p.budget(q, 0).values())
        headroom.append(dict(total_flow_m3h=q, each_filter_flow_m3h=q/2,
            assumed_nonfilter_loss_Pa=other,
            brochure_filter_headroom_Pa=p.fan(q)-other,
            numeric_filter_headroom_Pa=numeric(q)-other))
    tests = {
        'four_fans_add_flow_not_pressure': abs(numeric(4*xs[5]*p.CFM)-ys[5]*p.MMH2O)<1e-10,
        'two_filter_drops_not_added': p.budget(400,10)['filter_path']==10,
        'guard_area_R01': abs(.320**2*p.C['guard_open_fraction']-.036)<1e-12,
        'all_intersections_close': all(abs(r['closure_error_Pa'])<1e-8 for r in results),
        'numeric_nodes_recovered': all(abs(numeric(4*x*p.CFM)-y*p.MMH2O)<1e-10 for x,y in points),
        'loaded_flow_lower_both_curves': all(results[i]['total_flow_m3h']>results[i+1]['total_flow_m3h']>results[i+2]['total_flow_m3h'] for i in (0,3)),
        'historical_R01_reproduced': abs(results[0]['total_flow_m3h']-359.1739)<.001,
        'zero_losses_at_zero': sum(p.budget(0,20).values())==0,
        'numeric_endpoints': numeric(0)>0 and abs(numeric(4*xs[-1]*p.CFM))<1e-10,
    }
    assert all(tests.values()), tests
    result = dict(date='2026-10-03', status='REVIEW ONLY - ASSUMED FILTERS AND INSTALLED LOSSES',
        geometry='D01 R01 unchanged', source_url=URL,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        workbook_header_note=sheet['B12'].value,
        workbook_note='C=0.95 is retained as an unexplained OEM header annotation; no extra 0.95 correction applied.',
        numeric_points_count=len(points), rpm=2800,
        numeric_endpoints=dict(CFM=xs[-1],mmH2O=ys[0]),
        brochure_endpoints=dict(CFM=95,mmH2O=4.18),
        disposition='Do not average curves or treat either as guaranteed installed performance. OEM revision/test-basis explanation pending.',
        scenarios=results, filter_headroom=headroom, tests=tests)
    (OUT/'pressure_comparison.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    with (OUT/'OEM_P14_Max_points.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['single_fan_CFM','static_mmH2O','single_fan_m3h','static_Pa'])
        for x,y in points:w.writerow([x,y,x*p.CFM,y*p.MMH2O])
    with (OUT/'filter_headroom.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(headroom[0]));w.writeheader();w.writerows(headroom)
    with (OUT/'curve_comparison.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['total_flow_m3h','brochure_Pa','numeric_Pa','system_5Pa_ref','system_10Pa_ref','system_20Pa_ref'])
        for q in range(0,746,5):w.writerow([q,p.fan(q),numeric(q),*[sum(p.budget(q,r).values()) for r in (5,10,20)]])
    print(json.dumps(result,indent=2))

if __name__ == '__main__': main()
